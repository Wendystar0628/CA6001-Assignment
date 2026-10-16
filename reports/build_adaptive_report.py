"""Build the formal CA6001 report from its retained template and verified facts.

Run with python-docx 1.2.0 after results/adaptive_v1 is complete.
The fact file holds prose and table records; the Markdown is the editable text.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
REPORTS = ROOT / "reports"
REFERENCE = ROOT / "CA6001 Report Template.docx"
PRESERVE = ("word/theme/theme1.xml", "word/fontTable.xml", "word/numbering.xml", "word/webSettings.xml", "_rels/.rels")


def xml_child(parent, name, **attrs):
    child = OxmlElement(name)
    for key, value in attrs.items():
        child.set(qn(f"w:{key}"), str(value))
    parent.append(child)
    return child


def configure(doc):
    for name in ("Normal", "Title", "Heading 1", "Heading 2", "Caption"):
        if name not in doc.styles:
            doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.get_or_add_rPr().append(OxmlElement("w:rFonts")) if style._element.get_or_add_rPr().find(qn("w:rFonts")) is None else None
        for key in ("ascii", "hAnsi", "eastAsia", "cs"):
            style._element.get_or_add_rPr().find(qn("w:rFonts")).set(qn(f"w:{key}"), "Times New Roman")
        for node in style._element.xpath(".//w:pBdr"):
            node.getparent().remove(node)
    normal = doc.styles["Normal"]
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.05
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.widow_control = True
    for name, size in (("Heading 1", 14), ("Heading 2", 12)):
        s = doc.styles[name]
        s.font.size = Pt(size)
        s.font.bold = True
        s.paragraph_format.space_before = Pt(9 if name == "Heading 1" else 7)
        s.paragraph_format.space_after = Pt(5)
        s.paragraph_format.keep_with_next = True
    doc.styles["Title"].font.size = Pt(24)
    doc.styles["Title"].font.bold = True
    doc.styles["Caption"].font.size = Pt(10.5)
    doc.styles["Caption"].font.italic = False
    doc.styles["Caption"].paragraph_format.space_after = Pt(6)
    doc.styles["Caption"].paragraph_format.line_spacing = 1.0
    section = doc.sections[0]
    section.different_first_page_header_footer = True
    page_numbers = section._sectPr.find(qn("w:pgNumType"))
    if page_numbers is None:
        page_numbers = xml_child(section._sectPr, "w:pgNumType")
    page_numbers.set(qn("w:start"), "0")
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.font.name = "Times New Roman"
    run.font.size = Pt(10)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    run._r.addnext(field)
    settings = doc.settings._element
    field_setting = settings.find(qn("w:updateFields"))
    if field_setting is None:
        field_setting = xml_child(settings, "w:updateFields")
    field_setting.set(qn("w:val"), "true")


def text_runs(p, text):
    for part in re.split(r"(\*\*.*?\*\*|\*[^*]+\*)", text):
        bold = part.startswith("**") and part.endswith("**")
        italic = not bold and part.startswith("*") and part.endswith("*")
        content = part[2:-2] if bold else part[1:-1] if italic else part
        run = p.add_run(content)
        run.bold = bold
        run.italic = italic
        run.font.color.rgb = RGBColor(0, 0, 0)


def add_table(doc, table, number):
    p = doc.add_paragraph(f"Table {number}. {table['caption']}", "Caption")
    p.paragraph_format.keep_with_next = True
    headers, rows = table["columns"], table["rows"]
    widths = table.get("widths", [6.1 / len(headers)] * len(headers))
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for col, width in zip(t.columns, widths):
        col.width = Inches(width)
    pr = t._tbl.tblPr
    border = xml_child(pr, "w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        xml_child(border, f"w:{edge}", val="single", sz=4, color="D9D9D9")
    margins = xml_child(pr, "w:tblCellMar")
    for edge in ("top", "left", "bottom", "right"):
        xml_child(margins, f"w:{edge}", w=75, type="dxa")
    for i, header in enumerate(headers):
        t.rows[0].cells[i].text = str(header)
    xml_child(t.rows[0]._tr.get_or_add_trPr(), "w:tblHeader")
    for row in rows:
        for cell, value in zip(t.add_row().cells, row):
            cell.text = str(value)
    for row_i, row in enumerate(t.rows):
        for col_i, cell in enumerate(row.cells):
            cell.width = Inches(widths[col_i])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_i == 0:
                xml_child(cell._tc.get_or_add_tcPr(), "w:shd", fill="E6EDF2")
            for para in cell.paragraphs:
                para.alignment = WD_ALIGN_PARAGRAPH.LEFT if col_i == 0 else WD_ALIGN_PARAGRAPH.CENTER
                para.paragraph_format.space_after = Pt(1)
                para.paragraph_format.line_spacing = 1.0
                for run in para.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(table.get("font_pt", 10.5))
                    run.font.color.rgb = RGBColor(0, 0, 0)
                    run.bold = row_i == 0
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_figure(doc, figure, width, number, draft=False):
    image_path = ROOT / figure["path"]
    if not image_path.is_file():
        if draft:
            doc.add_paragraph(f"Figure {number} pending {figure.get('caption', '')}", "Caption")
            return
        raise FileNotFoundError(image_path)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(image_path), width=Inches(width))
    doc.add_paragraph(f"Figure {number}. {figure['caption']}", "Caption")


def visible_markdown(text, facts):
    """Expand data directives into portable Markdown tables and figure links."""
    figure_number = table_number = 0
    def table_block(match):
        nonlocal table_number
        table_number += 1
        key = match.group(1)
        table = facts.get("tables", {}).get(key)
        if table is None:
            return f"Table {table_number} pending {key}"
        def row(cells):
            return "| " + " | ".join(str(x).replace("|", "\\|") for x in cells) + " |"
        return "\n".join([f"Table {table_number}. {table['caption']}", "", row(table["columns"]), row(["---"] * len(table["columns"])), *map(row, table["rows"])])
    def figure_block(match):
        nonlocal figure_number
        figure_number += 1
        key = match.group(1)
        figure = facts.get("figures", {}).get(key)
        if figure is None:
            return f"Figure {figure_number} pending {key}"
        return f"![Figure {figure_number}. {figure['caption']}](../{figure['path']})"
    text = re.sub(r"<!-- TABLE ([a-z_]+) -->", table_block, text)
    return re.sub(r"<!-- FIGURE ([a-z_]+) [0-9.]+ -->", figure_block, text)


def build(facts_path, draft=False):
    facts = json.loads(facts_path.read_text(encoding="utf-8"))
    facts["text"]["references"] = (REPORTS / "references.md").read_text(encoding="utf-8").strip()
    if not draft and facts.get("status") != "final":
        raise ValueError("Final report requires final verified report_facts.json")
    template_hash = hashlib.sha256(REFERENCE.read_bytes()).hexdigest()
    text = (REPORTS / "report_template.md").read_text(encoding="utf-8")
    missing = []
    def substitute(match):
        key = match.group(1)
        if key not in facts["text"]:
            missing.append(key)
            return f"PENDING {key}" if draft else match.group(0)
        return str(facts["text"][key])
    text = re.sub(r"\{\{([a-z_]+)\}\}", substitute, text)
    if missing and not draft:
        raise ValueError(f"Missing verified report text: {missing}")
    doc = Document(REFERENCE)
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)
    configure(doc)
    doc.core_properties.author = "CA6001 Project Team"
    doc.core_properties.title = "Machine Learning Guided Adaptive CPT Investigation and Economic Benefit Assessment"
    figure_number = table_number = 0
    cover = True
    references = False
    page_break_before = False
    for block in re.split(r"\n\s*\n", text):
        block = block.strip()
        if not block:
            continue
        if block == "<!-- PAGE -->":
            page_break_before = True
            cover = False
            continue
        match = re.fullmatch(r"<!-- TABLE ([a-z_]+) -->", block)
        if match:
            key = match.group(1)
            table_number += 1
            if key in facts.get("tables", {}):
                add_table(doc, facts["tables"][key], table_number)
            elif draft:
                doc.add_paragraph(f"Table {table_number} pending {key}", "Caption")
            else:
                raise ValueError(f"Missing table {key}")
            continue
        match = re.fullmatch(r"<!-- FIGURE ([a-z_]+) ([0-9.]+) -->", block)
        if match:
            key, width = match.group(1), float(match.group(2))
            figure_number += 1
            figure = facts.get("figures", {}).get(key, {"path": f"results/adaptive_v1/figures/{key}.png", "caption": key})
            add_figure(doc, figure, width, figure_number, draft)
            continue
        if block.startswith("### "):
            p = doc.add_paragraph(block[4:], "Heading 2")
        elif block.startswith("## "):
            title = block[3:]
            references = title.startswith("10 References")
            p = doc.add_paragraph(title, "Heading 1")
        elif block.startswith("# "):
            p = doc.add_paragraph(block[2:], "Title")
            p.paragraph_format.space_before = Pt(60)
            p.paragraph_format.space_after = Pt(28)
        else:
            p = doc.add_paragraph()
            text_runs(p, block)
            if cover:
                p.paragraph_format.space_after = Pt(18)
            elif references:
                p.paragraph_format.left_indent = Inches(0.22)
                p.paragraph_format.first_line_indent = Inches(-0.22)
                p.paragraph_format.space_after = Pt(5)
                p.paragraph_format.line_spacing = 1.0
                for run in p.runs:
                    run.font.size = Pt(10.5)
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if page_break_before:
            p.paragraph_format.page_break_before = True
            page_break_before = False
    out = REPORTS / (".qa/CA6001_Adaptive_CPT_Report_draft.docx" if draft else "CA6001_Adaptive_CPT_Report.docx")
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    # python-docx may serialise untouched package parts; restore preserve-only parts.
    with ZipFile(REFERENCE) as z:
        preserved = {name: z.read(name) for name in PRESERVE}
    with ZipFile(out) as z:
        package = {name: z.read(name) for name in z.namelist()}
    package.update(preserved)
    with ZipFile(out, "w", ZIP_DEFLATED) as z:
        for name, payload in package.items():
            z.writestr(name, payload)
    md_out = REPORTS / (".qa/draft_report.md" if draft else "CA6001_Adaptive_CPT_Report.md")
    md_out.write_text(visible_markdown(text, facts), encoding="utf-8")
    with ZipFile(out) as z:
        assert all(z.read(name) == payload for name, payload in preserved.items())
    assert hashlib.sha256(REFERENCE.read_bytes()).hexdigest() == template_hash
    receipt = {"template_sha256": template_hash, "preserved_parts": list(preserved), "figures": figure_number, "tables": table_number, "missing_text": missing, "evidence": facts.get("evidence", [])}
    (REPORTS / ".qa/build_receipt.json").write_text(json.dumps(receipt, ensure_ascii=False, indent=2), encoding="utf-8")
    print(out)


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser()
    parser.add_argument("--facts", type=Path, default=REPORTS / "report_facts.json")
    parser.add_argument("--draft", action="store_true")
    args = parser.parse_args()
    build(args.facts, args.draft)
